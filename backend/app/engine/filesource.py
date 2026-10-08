"""Files as a source: CSV/TSV, JSON/NDJSON, Parquet (DuckDB) and Excel (openpyxl).

Implements the read side of the Adapter interface, so the runner treats a file like a table. CSV columns are
emitted as text and cast by the transform step, which is what gives row-level "cannot convert" rejects.
Rows the CSV parser itself cannot make sense of (wrong column count, broken quoting) come back as rejects too.
"""
from __future__ import annotations

from collections.abc import Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import duckdb
import pyarrow as pa

from ..connectors.base import UserFacingError
from ..connectors.paths import resolve_allowed
from ..files import detect_format, materialize, split_compression
from .arrowconv import conform, rows_to_table_lenient
from .ddl import qi
from .model import Batch, ReadRequest, SplitSpec
from .types import ColSpec, normalize

_EXACT_COUNT_MAX_BYTES = 200 * 1024 * 1024


@dataclass
class FileSpec:
    path: str
    format: str | None = None  # csv | json | parquet | excel; detected from the extension when empty
    delimiter: str | None = None  # csv; auto-detected when empty
    header: bool = True
    encoding: str = "utf-8"
    sheet: str | None = None  # excel; first sheet when empty
    skip_rows: int = 0
    null_values: list[str] = field(default_factory=lambda: [""])  # csv strings that mean NULL


def _q(s: str) -> str:
    return "'" + s.replace("'", "''") + "'"


class FileAdapter:
    """Source-only adapter over one file."""

    dialect = "file"

    def __init__(self, spec: FileSpec, *, trusted: Path | None = None, label: str | None = None,
                 unpack_dir: Path | None = None):
        """`trusted`: a local copy Conduit made itself (a remote download), read without the allowed-roots check.
        `unpack_dir`: where a .zip's data file is unpacked (a job's own folder); default: the shared file cache."""
        self.spec = spec
        self.source: Path = trusted if trusted is not None else resolve_allowed(spec.path)  # what the user named
        if not self.source.exists() or not self.source.is_file():
            raise FileNotFoundError(f"{self.source} does not exist")
        self.inner_name, self.compression = split_compression(self.source.name)
        self.format = spec.format or detect_format(self.source)
        if self.inner_name.lower().endswith(".xls") and self.format == "excel":
            raise ValueError("Legacy .xls files are not supported: save the sheet as .xlsx or .csv")
        self.path: Path = materialize(self.source, unpack_dir)  # what is actually read (a .zip's data file, unpacked)
        self.label = label or self.source.name
        self.notes: list[str] = []
        self.max_writers = 0

    # ---- DuckDB relation ---------------------------------------------------------
    def _relation(self, all_text: bool = True) -> str:
        p, s = _q(str(self.path)), self.spec
        if self.format == "csv":
            opts = [f"header={'true' if s.header else 'false'}", f"encoding={_q(s.encoding)}", "ignore_errors=true",
                    "store_rejects=true", f"nullstr=[{', '.join(_q(v) for v in s.null_values)}]", "null_padding=false"]
            if all_text:
                opts.append("all_varchar=true")
            delim = s.delimiter or ("\t" if self.inner_name.lower().endswith(".tsv") or self.path.suffix.lower() == ".tsv" else None)
            if delim:
                opts.append(f"delim={_q(delim)}")
            if s.skip_rows:
                opts.append(f"skip={int(s.skip_rows)}")
            return f"read_csv({p}, {', '.join(opts)})"
        if self.format == "json":
            return f"read_json_auto({p})"
        return f"read_parquet({p})"

    # ---- columns -----------------------------------------------------------------
    def read_columns(self) -> list[ColSpec]:
        """Columns exactly as batches will carry them (CSV: all text)."""
        if self.format == "excel":
            return self._excel_columns()[0]
        con = duckdb.connect()
        try:
            rel = con.sql(f"SELECT * FROM {self._relation()} LIMIT 0")
            return [ColSpec(n, normalize("duckdb", str(t)), True) for n, t in zip(rel.columns, rel.types)]
        finally:
            con.close()

    def suggest_columns(self) -> list[ColSpec]:
        """Best-guess typed columns, used as the wizard's default target types."""
        if self.format == "excel":
            return self._excel_columns()[0]
        con = duckdb.connect()
        try:
            opts = self._relation(all_text=False)
            rel = con.sql(f"SELECT * FROM {opts} LIMIT 0")
            return [ColSpec(n, normalize("duckdb", str(t)), True) for n, t in zip(rel.columns, rel.types)]
        finally:
            con.close()

    def _excel_columns(self) -> tuple[list[ColSpec], list[str]]:
        from openpyxl import load_workbook

        wb = load_workbook(self.path, read_only=True, data_only=True)
        try:
            ws = wb[self.spec.sheet] if self.spec.sheet else wb.worksheets[0]
            rows = ws.iter_rows(values_only=True)
            for _ in range(self.spec.skip_rows):
                next(rows, None)
            first = next(rows, None)
            if first is None:
                return [], []
            width = len(first)
            names: list[str] = []
            for i, h in enumerate(first if self.spec.header else [None] * width):
                base = str(h).strip() if h not in (None, "") else f"col_{i + 1}"
                name, k = base, 1
                while name in names:
                    k += 1
                    name = f"{base}_{k}"
                names.append(name)
            seen: list[set[str]] = [set() for _ in names]
            sample = [first] if not self.spec.header else []
            for i, row in enumerate(rows):
                if i >= 1000:
                    break
                sample.append(row)
            for row in sample:
                for j, v in enumerate(row[:width]):
                    if v is not None:
                        seen[j].add(type(v).__name__)
            def logical(ts: set[str]) -> str:
                if not ts:
                    return "VARCHAR"
                if ts <= {"int"}:
                    return "BIGINT"
                if ts <= {"int", "float"}:
                    return "DOUBLE"
                if ts == {"bool"}:
                    return "BOOLEAN"
                if ts == {"datetime"}:
                    return "TIMESTAMP"
                if ts == {"date"}:
                    return "DATE"
                return "VARCHAR"
            return [ColSpec(n, logical(t), True) for n, t in zip(names, seen)], names
        finally:
            wb.close()

    # ---- Adapter-like interface --------------------------------------------------
    def estimate_rows(self, req: ReadRequest | None = None, exact: bool = False, cheap: bool = False) -> tuple[int | None, bool]:
        """`cheap`: never count a text file exactly (many files in one job: estimate each instead of scanning all)."""
        size = self.path.stat().st_size
        if self.format == "excel":
            from openpyxl import load_workbook

            wb = load_workbook(self.path, read_only=True, data_only=True)
            try:
                ws = wb[self.spec.sheet] if self.spec.sheet else wb.worksheets[0]
                return max((ws.max_row or 1) - (1 if self.spec.header else 0) - self.spec.skip_rows, 0), True
            finally:
                wb.close()
        gz = self.compression == "gzip"
        con = duckdb.connect()
        try:
            if self.format == "parquet" or exact or (not cheap and size <= _EXACT_COUNT_MAX_BYTES // (5 if gz else 1)):
                return int(con.sql(f"SELECT count(*) FROM {self._relation()}").fetchone()[0]), True
            if gz:  # compressed bytes say little about the row count: counted during the load instead
                return None, False
            sample = con.sql(f"SELECT count(*) FROM (SELECT * FROM {self._relation()} LIMIT 20000)").fetchone()[0]
            with self.path.open("rb") as f:
                head = f.read(min(size, 8 * 1024 * 1024))
            lines = max(head.count(b"\n"), 1)
            return int(size / (len(head) / lines)) if sample else 0, False
        finally:
            con.close()

    def key_range(self, req: ReadRequest) -> None:
        return None

    def plan_splits(self, req: ReadRequest, n: int) -> list[SplitSpec]:
        # One reader per file (DuckDB parallelises inside it). A "file" split resumes by row offset: the file is a
        # fixed snapshot, so a retry skips exactly the rows already committed (in file order).
        return [SplitSpec(0, self.label, "file", resumable=True)]

    def read_batches(self, req: ReadRequest, split: SplitSpec, batch_rows: int) -> Iterator[Batch]:
        """Batches in file order. Each carries the number of rows read so far (the split's resume offset)."""
        offset = int(split.extra.get("offset") or 0)
        if self.format == "excel":
            yield from self._read_excel(req, batch_rows, offset)
            return
        cols = ", ".join(qi("duckdb", c.name) for c in req.columns)
        con = duckdb.connect()
        try:
            con.execute("SET preserve_insertion_order = true")  # file order, so OFFSET means "the first n rows"
            res = con.execute(f"SELECT {cols} FROM {self._relation()}" + (f" OFFSET {offset}" if offset else ""))
            reader = (res.to_arrow_reader if hasattr(res, "to_arrow_reader") else res.fetch_record_batch)(batch_rows)
            done = offset
            for rb in reader:
                if rb.num_rows:
                    done += rb.num_rows
                    yield Batch(conform(pa.Table.from_batches([rb]), req.columns), done)
            if self.format == "csv":
                yield Batch(pa.table({c.name: pa.array([], type=pa.string()) for c in req.columns}), None, self._csv_rejects(con, req))
        finally:
            con.close()

    def _csv_rejects(self, con: duckdb.DuckDBPyConnection, req: ReadRequest) -> list[tuple[str, dict[str, Any]]]:
        try:
            rows = con.execute("SELECT line, column_name, error_type, error_message, csv_line FROM reject_errors ORDER BY line").fetchall()
        except duckdb.Error:
            return []  # no rejects table: nothing was rejected
        by_line: dict[int, tuple[list[str], str]] = {}
        for ln, col, etype, _msg, raw in rows:  # one reject per bad line, even when several columns are affected
            reasons, _ = by_line.setdefault(ln, ([], (raw or "").strip("\r\n")))
            label = f"{etype} ({col})" if col else str(etype)
            if label not in reasons:
                reasons.append(label)
        return [(f"line {ln}: " + "; ".join(r), {"line": ln, "raw": raw}) for ln, (r, raw) in sorted(by_line.items())]

    def _read_excel(self, req: ReadRequest, batch_rows: int, offset: int = 0) -> Iterator[Batch]:
        from openpyxl import load_workbook

        cols, names = self._excel_columns()
        idx = {n: i for i, n in enumerate(names)}
        wanted = [idx[c.name] for c in req.columns]
        wb = load_workbook(self.path, read_only=True, data_only=True)
        try:
            ws = wb[self.spec.sheet] if self.spec.sheet else wb.worksheets[0]
            rows = ws.iter_rows(values_only=True)
            for _ in range(self.spec.skip_rows + (1 if self.spec.header else 0)):
                next(rows, None)
            buf: list[tuple] = []
            seen = 0  # data rows so far (blank rows are not rows); the first `offset` were loaded before a retry
            for row in rows:
                if all(v is None for v in row):
                    continue
                seen += 1
                if seen <= offset:
                    continue
                padded = tuple(row) + (None,) * (len(names) - len(row))
                buf.append(tuple(padded[i] for i in wanted))
                if len(buf) >= batch_rows:
                    table, rejects = rows_to_table_lenient(buf, req.columns)
                    yield Batch(table, seen, rejects)
                    buf = []
            if buf:
                table, rejects = rows_to_table_lenient(buf, req.columns)
                yield Batch(table, seen, rejects)
        finally:
            wb.close()


# ---- several files as one source -------------------------------------------------------------

class FileSetError(UserFacingError):
    """The files of one source do not line up (different columns)."""


@dataclass
class SourceFile:
    label: str  # how the file is named in splits, rejects and notes (never carries a URL query string)
    adapter: FileAdapter | None  # None while planning: only the first file is fetched for a preview
    size: int | None = None
    remote: dict[str, Any] | None = None  # the manifest record of a downloaded remote file
    est_rows: int | None = None


class MultiFileAdapter:
    """One or more files with the same column names read as one source: one split per file, in name order.

    Each split is a "file" split that resumes by row offset, so a retry continues a half-loaded file exactly and
    starts the files not done yet. Columns come from the first file; every other one must have the same names
    (any order), checked by check_columns() before anything is written.
    """

    dialect = "file"
    max_writers = 0

    def __init__(self, files: list[SourceFile], notes: list[str] | None = None) -> None:
        if not files or files[0].adapter is None:
            raise ValueError("MultiFileAdapter needs at least its first file")
        self.files = files
        self.first = files[0].adapter
        self.format = self.first.format
        self.notes = list(notes or [])

    def read_columns(self) -> list[ColSpec]:
        return self.first.read_columns()

    def suggest_columns(self) -> list[ColSpec]:
        return self.first.suggest_columns()

    def check_columns(self) -> None:
        want = {c.name for c in self.first.read_columns()}
        for f in self.files[1:]:
            if f.adapter is None:
                continue
            have = {c.name for c in f.adapter.read_columns()}
            if have != want:
                parts = []
                if missing := sorted(want - have):
                    parts.append(f"missing {', '.join(missing[:8])}")
                if extra := sorted(have - want):
                    parts.append(f"extra {', '.join(extra[:8])}")
                raise FileSetError(f"{f.label} does not have the same columns as {self.files[0].label} "
                                   f"({'; '.join(parts)}). All files of one load need the same column names.")

    def estimate_rows(self, req: ReadRequest | None = None, exact: bool = False) -> tuple[int | None, bool]:
        if any(f.adapter is None for f in self.files):
            return None, False  # only the first file is known while planning: counted during the load
        sizes = sum((f.adapter.path.stat().st_size for f in self.files), 0)
        cheap = len(self.files) > 1 and sizes > _EXACT_COUNT_MAX_BYTES  # don't scan gigabytes up front
        total, all_exact = 0, True
        for f in self.files:
            n, ex = f.adapter.estimate_rows(exact=exact, cheap=cheap)
            if n is None:
                return None, False
            f.est_rows = n
            total += n
            all_exact = all_exact and ex
        return total, all_exact

    def key_range(self, req: ReadRequest) -> None:
        return None

    def plan_splits(self, req: ReadRequest, n: int) -> list[SplitSpec]:
        return [SplitSpec(i, f.label, "file", est_rows=f.est_rows, resumable=True, extra={"file": i})
                for i, f in enumerate(self.files)]

    def reject_prefix(self, split: SplitSpec) -> str:
        """Prefix for reject reasons, so a reject says which file it came from (only when there are several)."""
        return f"{self.files[split.extra.get('file', split.idx)].label}: " if len(self.files) > 1 else ""

    def read_batches(self, req: ReadRequest, split: SplitSpec, batch_rows: int) -> Iterator[Batch]:
        f = self.files[split.extra.get("file", split.idx)]
        if f.adapter is None:
            raise FileSetError(f"{f.label} has not been fetched")
        prefix = self.reject_prefix(split)
        for b in f.adapter.read_batches(req, split, batch_rows):
            if prefix and b.rejects:
                b.rejects = [(prefix + reason, row) for reason, row in b.rejects]
            yield b
