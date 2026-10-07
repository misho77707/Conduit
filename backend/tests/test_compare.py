"""Compare: source vs target diffs end to end through the API, on SQLite / DuckDB files (no live server needed)."""
import csv
import io
import json
import sqlite3
import time
import zipfile

import duckdb
import pyarrow.parquet as pq
import pytest

from jobhelpers import add_conn, uniq


def _sqlite(env, name, ddl, rows):
    p = env / name
    c = sqlite3.connect(p)
    c.execute(ddl)
    if rows:
        c.executemany(f"insert into {ddl.split()[2].split('(')[0]} values ({', '.join('?' * len(rows[0]))})", rows)
    c.commit()
    c.close()
    return p


def _conn(client, auth, type_, path):
    return add_conn(client, auth, type_, uniq(type_), {"path": str(path), "create_if_missing": True})


def _side(conn_id, table, **kw):
    return {"connection_id": conn_id, "table": table, **kw}


def _run(client, auth, body, timeout=60):
    r = client.post("/api/compare", json=body, headers=auth)
    assert r.status_code == 201, r.text
    cid = r.json()["id"]
    end = time.time() + timeout
    while time.time() < end:
        c = client.get(f"/api/compare/{cid}", headers=auth).json()
        if c["status"] in ("succeeded", "failed", "cancelled"):
            return c
        time.sleep(0.1)
    raise AssertionError(f"comparison did not finish: {c}")


SRC_ROWS = [(1, "a", 10.0), (2, "b", 20.0), (3, "c", 30.0), (4, "d", 40.0), (5, "e", 50.0), (6, "f", 60.0)]
TGT_ROWS = [(1, "a", 10.0), (2, "b", 21.0), (4, None, 40.0), (5, "e", 50.0), (6, "f  ", 60.0), (7, "g", 70.0)]


@pytest.fixture()
def pair(client, auth, env):
    s = _conn(client, auth, "sqlite", _sqlite(env, "s.sqlite", "create table src(id integer primary key, name text, amount real)", SRC_ROWS))
    t = _conn(client, auth, "sqlite", _sqlite(env, "t.sqlite", "create table tgt(id integer primary key, name text, amount real)", TGT_ROWS))
    return s, t


def test_columns_endpoint_pairs_names_and_suggests_the_primary_key(client, auth, pair):
    s, t = pair
    r = client.post("/api/compare/columns", json={"source": _side(s, "src"), "target": _side(t, "tgt")}, headers=auth)
    assert r.status_code == 200, r.text
    body = r.json()
    assert [p["target"] for p in body["pairs"]] == ["id", "name", "amount"]
    assert body["suggested_keys"] == ["id"] and "primary key" in body["key_note"]


def test_key_mode_finds_missing_extra_and_changed_rows(client, auth, pair):
    s, t = pair
    c = _run(client, auth, {"source": _side(s, "src"), "target": _side(t, "tgt")})
    assert c["status"] == "succeeded", c
    sm = c["summary"]
    assert sm["mode"] == "key" and sm["keys"] == ["id"]
    assert (sm["rows_source"], sm["rows_target"]) == (6, 6)
    assert (sm["missing"], sm["extra"], sm["mismatched"], sm["matched"]) == (1, 1, 3, 2)
    assert sm["column_mismatches"] == {"name": 2, "amount": 1}
    assert sm["identical"] is False

    missing = client.get(f"/api/compare/{c['id']}/rows", params={"set": "missing"}, headers=auth).json()
    assert missing["rows"] == [{"id": 3, "name": "c", "amount": 30.0}]
    extra = client.get(f"/api/compare/{c['id']}/rows", params={"set": "extra"}, headers=auth).json()
    assert extra["rows"][0]["id"] == 7
    mm = client.get(f"/api/compare/{c['id']}/rows", params={"set": "mismatched"}, headers=auth).json()
    assert mm["columns"] == ["id", "name__source", "name__target", "amount__source", "amount__target", "_diff_columns"]
    by_id = {r["id"]: r for r in mm["rows"]}
    assert by_id[2]["_diff_columns"] == "amount" and (by_id[2]["amount__source"], by_id[2]["amount__target"]) == (20.0, 21.0)
    assert by_id[4]["_diff_columns"] == "name" and by_id[4]["name__target"] is None
    assert by_id[6]["name__target"] == "f  "


def test_ignore_nulls_and_trim_text(client, auth, pair):
    s, t = pair
    c = _run(client, auth, {"source": _side(s, "src"), "target": _side(t, "tgt"),
                            "options": {"ignore_nulls": True, "trim_text": True}})
    assert c["summary"]["mismatched"] == 1 and c["summary"]["column_mismatches"] == {"amount": 1}


def test_whole_row_mode_counts_duplicates(client, auth, env):
    s = _conn(client, auth, "sqlite", _sqlite(env, "a.sqlite", "create table r(k text, v integer)", [("A", 1), ("A", 1), ("B", 2)]))
    t = _conn(client, auth, "sqlite", _sqlite(env, "b.sqlite", "create table r(k text, v integer)", [("A", 1), ("B", 2), ("C", 3)]))
    c = _run(client, auth, {"source": _side(s, "r"), "target": _side(t, "r"), "keys": []})
    sm = c["summary"]
    assert sm["mode"] == "row" and (sm["missing"], sm["extra"], sm["matched"]) == (1, 1, 2)
    assert "mismatched" not in sm["sets"]
    dup = client.get(f"/api/compare/{c['id']}/rows", params={"set": "dup_source"}, headers=auth).json()
    assert dup["rows"] == [{"k": "A", "v": 1, "_count": 2}]
    r = client.get(f"/api/compare/{c['id']}/rows", params={"set": "mismatched"}, headers=auth)
    assert r.status_code == 422


def test_duplicate_keys_are_reported(client, auth, env):
    s = _conn(client, auth, "sqlite", _sqlite(env, "a.sqlite", "create table r(id integer, v text)", [(1, "x"), (2, "y")]))
    t = _conn(client, auth, "sqlite", _sqlite(env, "b.sqlite", "create table r(id integer, v text)", [(1, "x"), (2, "y"), (2, "y")]))
    c = _run(client, auth, {"source": _side(s, "r"), "target": _side(t, "r"), "keys": ["id"]})
    sm = c["summary"]
    assert (sm["dup_source_keys"], sm["dup_target_keys"], sm["dup_target_rows"]) == (0, 1, 2)
    assert any("more than once" in n for n in sm["notes"])


def test_types_are_compared_by_value_across_databases(client, auth, env):
    s = _conn(client, auth, "sqlite", _sqlite(env, "a.sqlite", "create table r(id text, price text, ratio real, label text)",
                                              [("1", " 1.50", 0.1 + 0.2, "x"), ("2", "abc", 1.0, "y")]))
    d = env / "t.duckdb"
    con = duckdb.connect(str(d))
    con.execute("create table r(ident integer, price decimal(10,2), ratio double, label varchar)")
    con.execute("insert into r values (1, 1.5, 0.3, 'x'), (2, 7, 1.0, 'y')")
    con.execute("create table z(id integer, happened timestamptz)")
    con.execute("insert into z values (1, TIMESTAMPTZ '2024-01-01 10:00:00+04')")
    con.close()
    t = _conn(client, auth, "duckdb", d)
    cols = [{"source": "id", "target": "ident"}, {"source": "price", "target": "price"}, {"source": "ratio", "target": "ratio"}]
    c = _run(client, auth, {"source": _side(s, "r"), "target": _side(t, "main.r"), "columns": cols, "keys": ["ident"]})
    sm = c["summary"]
    assert c["status"] == "succeeded", c
    assert [col["compared_as"] for col in sm["columns"]] == ["BIGINT", "DECIMAL(38,2)", "DOUBLE"]
    assert (sm["mismatched"], sm["column_mismatches"]) == (1, {"price": 1})  # "abc" is not a number: NULL vs 7

    s2 = _conn(client, auth, "sqlite", _sqlite(env, "z.sqlite", "create table z(id integer, happened timestamp)", [(1, "2024-01-01 06:00:00")]))
    c = _run(client, auth, {"source": _side(s2, "z"), "target": _side(t, "main.z"), "keys": ["id"]})
    assert c["summary"]["identical"] is True, c["summary"]


def test_identical_tables(client, auth, env):
    p = _sqlite(env, "same.sqlite", "create table r(id integer primary key, v text)", [(1, "x"), (2, None)])
    s = _conn(client, auth, "sqlite", p)
    c = _run(client, auth, {"source": _side(s, "r"), "target": _side(s, "r")})
    assert c["summary"]["identical"] is True and c["summary"]["matched"] == 2


def test_file_source_against_a_table(client, auth, env, pair):
    _, t = pair
    (env / "src.csv").write_text("id,name,amount\n1,a,10\n2,b,21\n4,,40\n5,e,50\n6,f  ,60\n7,g,70\n")
    c = _run(client, auth, {"source": {"file": {"path": str(env / "src.csv")}}, "target": _side(t, "tgt"), "keys": ["id"]})
    assert c["status"] == "succeeded", c
    assert c["summary"]["identical"] is True, c["summary"]


def test_exports(client, auth, pair):
    s, t = pair
    c = _run(client, auth, {"source": _side(s, "src"), "target": _side(t, "tgt")})
    url = f"/api/compare/{c['id']}/export"

    r = client.get(url, params={"sets": "missing", "format": "csv"}, headers=auth)
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert list(csv.reader(io.StringIO(r.text))) == [["id", "name", "amount"], ["3", "c", "30.0"]]

    r = client.get(url, params={"sets": "mismatched", "format": "csv", "columns": "amount"}, headers=auth)
    assert next(csv.reader(io.StringIO(r.text))) == ["id", "amount__source", "amount__target", "_diff_columns"]

    r = client.get(url, params={"sets": "missing,extra,mismatched", "format": "parquet"}, headers=auth)
    zf = zipfile.ZipFile(io.BytesIO(r.content))
    names = sorted(zf.namelist())
    assert len(names) == 3 and all(n.endswith(".parquet") for n in names)
    assert pq.read_table(io.BytesIO(zf.read([n for n in names if n.endswith("extra.parquet")][0]))).num_rows == 1

    r = client.get(url, params={"sets": "mismatched", "format": "json", "rows": "sample", "columns": "keys"}, headers=auth)
    lines = [json.loads(x) for x in r.text.splitlines()]
    assert len(lines) == 3 and set(lines[0]) == {"id", "_diff_columns"}

    from openpyxl import load_workbook
    r = client.get(url, params={"sets": "missing,extra", "format": "xlsx"}, headers=auth)
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["missing in target", "extra in target"]
    assert [c.value for c in wb["extra in target"][2]] == [7, "g", 70.0]

    assert client.get(url, params={"sets": "nope", "format": "csv"}, headers=auth).status_code == 422
    assert client.get(url, params={"sets": "missing", "format": "doc"}, headers=auth).status_code == 422
    assert client.get(url, params={"sets": "missing", "columns": "zzz"}, headers=auth).status_code == 422


def test_bad_requests_fail_fast(client, auth, pair):
    s, t = pair
    r = client.post("/api/compare", json={"source": _side(s, "src"), "target": _side(t, "nope")}, headers=auth)
    assert r.status_code == 422 and "nope" in r.text
    r = client.post("/api/compare", json={"source": _side(s, "src"), "target": _side(t, "tgt"), "keys": ["zz"]}, headers=auth)
    assert r.status_code == 422 and "zz" in r.text
    r = client.post("/api/compare", json={"source": _side(s, "src"), "target": _side(t, "tgt"),
                                          "columns": [{"source": "id", "target": "id", "include": False}]}, headers=auth)
    assert r.status_code == 422 and "No columns" in r.text


def test_only_the_latest_comparisons_are_kept(client, auth, pair, monkeypatch):
    from app.config import get_settings
    from app.compare.engine import folder

    monkeypatch.setenv("CONDUIT_KEEP_COMPARISONS", "2")
    get_settings.cache_clear()
    s, t = pair
    ids = [_run(client, auth, {"source": _side(s, "src"), "target": _side(t, "tgt")})["id"] for _ in range(3)]
    listed = [c["id"] for c in client.get("/api/compare", headers=auth).json()]
    assert listed == ids[:0:-1]  # the two newest, newest first
    assert not folder(ids[0]).exists() and folder(ids[2]).exists()


def test_cancel_and_delete(client, auth, pair):
    from app.db import get_sessionmaker
    from app.models import Comparison

    with get_sessionmaker()() as s:
        c = Comparison(name="queued one", spec={}, status="queued")
        s.add(c)
        s.commit()
        cid = str(c.id)
    assert client.delete(f"/api/compare/{cid}", headers=auth).status_code == 409
    r = client.post(f"/api/compare/{cid}/cancel", headers=auth)
    assert r.status_code == 200 and r.json()["status"] == "cancelled"
    assert client.delete(f"/api/compare/{cid}", headers=auth).status_code == 204
    assert client.get(f"/api/compare/{cid}", headers=auth).status_code == 404
